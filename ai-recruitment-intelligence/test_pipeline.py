import os
import requests
import docx
from pypdf import PdfWriter

# Ensure test fixtures directory exists
TEST_DIR = os.path.join(os.path.dirname(__file__), "test_fixtures")
os.makedirs(TEST_DIR, exist_ok=True)

def create_sample_files():
    # 1. Create a Sample JD (DOCX)
    jd_doc = docx.Document()
    jd_doc.add_heading("Senior Backend Engineer (Python / Flask)", level=1)
    jd_doc.add_paragraph("Company: Nexus Technologies")
    jd_doc.add_paragraph("Experience Required: Minimum 3+ years of professional backend engineering.")
    jd_doc.add_paragraph("Must-Have Requirements:\n- Strong proficiency in Python and Flask or FastAPI\n- Solid understanding of SQL and Relational Databases (PostgreSQL / SQLite)\n- Experience designing and securing RESTful APIs\n- Practical knowledge of Docker and Git version control")
    jd_doc.add_paragraph("Nice to Have:\n- Familiarity with Redis caching and Celery task queues\n- Basic AWS cloud deployment")
    jd_doc.add_paragraph("Responsibilities:\n- Build, test, and maintain robust web APIs\n- Collaborate with frontend engineers to ship features\n- Write unit tests and maintain high code quality")
    jd_path = os.path.join(TEST_DIR, "Sample_Backend_JD.docx")
    jd_doc.save(jd_path)

    # 2. Create Candidate A (Strong Fit) - TXT
    cand_a_text = """
Alex Morgan
Email: alex.morgan@email.com | Phone: (555) 234-5678
Location: San Francisco, CA

Professional Summary:
Senior Python Backend Developer with 4 years of experience building high-traffic REST APIs and microservices using Flask and PostgreSQL.

Technical Skills:
Python, Flask, FastAPI, PostgreSQL, SQL, Docker, Git, Redis, REST API

Work Experience:
Lead Backend Developer | CloudScale Systems (2022 - Present)
- Architected and deployed 15+ RESTful API microservices using Python and Flask serving 500k monthly requests.
- Optimized PostgreSQL database schema and complex SQL queries, reducing average API response times by 35%.
- Containerized development and staging environments using Docker, cutting onboarding time in half.
- Integrated Redis caching layer for session management and rate limiting.

Software Engineer | AppForge Inc. (2020 - 2022)
- Developed secure authentication endpoints and CRUD operations in Python.
- Managed Git workflow and automated testing pipelines.

Education:
B.S. in Computer Science | University of California (2020)
"""
    cand_a_path = os.path.join(TEST_DIR, "Alex_Morgan_Resume.txt")
    with open(cand_a_path, "w", encoding="utf-8") as f:
        f.write(cand_a_text.strip())

    # 3. Create Candidate B (Keyword Stuffer / False Positive) - DOCX
    cand_b_doc = docx.Document()
    cand_b_doc.add_heading("Jordan Taylor - Software Engineer", level=1)
    cand_b_doc.add_paragraph("Contact: jordan.taylor@email.com | Phone: (555) 987-6543")
    cand_b_doc.add_paragraph("SKILLS: Python, Flask, FastAPI, PostgreSQL, Docker, Kubernetes, AWS, GraphQL, Redis, Kafka, Machine Learning, PyTorch, React, Java, C++, Blockchain, CyberSecurity")
    cand_b_doc.add_paragraph("EXPERIENCE:")
    cand_b_doc.add_paragraph("Junior Web Assistant | WebStudio (2023 - 2024)\n- Updated WordPress blog themes and HTML landing pages.\n- Responded to client customer support emails.\n- Created weekly social media graphic banners.")
    cand_b_doc.add_paragraph("EDUCATION:\nB.A. in Digital Arts (2023)")
    cand_b_path = os.path.join(TEST_DIR, "Jordan_Taylor_Resume.docx")
    cand_b_doc.save(cand_b_path)

    # 4. Create Candidate C (Junior / Partial Fit) - TXT
    cand_c_text = """
Samantha Ray
Email: sam.ray@email.com | Phone: (555) 444-1122

Summary:
Junior Python Programmer with 1 year of experience building script utilities and simple web pages.

Skills:
Python, HTML, CSS, JavaScript, Git, SQLite

Experience:
Intern Developer | ByteLogic (2023 - 2024)
- Wrote automated Python scripts for data scraping and CSV cleaning.
- Maintained internal SQLite database records.

Education:
Associate Degree in Information Technology (2023)
"""
    cand_c_path = os.path.join(TEST_DIR, "Samantha_Ray_Resume.txt")
    with open(cand_c_path, "w", encoding="utf-8") as f:
        f.write(cand_c_text.strip())

    return jd_path, [cand_a_path, cand_b_path, cand_c_path]


def run_test():
    base_url = "http://127.0.0.1:5000"
    
    print("[1] Testing Health Endpoint...")
    h_resp = requests.get(f"{base_url}/api/health")
    print(f"    Health Status: {h_resp.status_code}, {h_resp.json()}")

    print("[2] Creating Sample JD & Resumes...")
    jd_path, resume_paths = create_sample_files()

    print("[3] Testing /api/upload endpoint...")
    with open(jd_path, "rb") as jd_f, \
         open(resume_paths[0], "rb") as r1, \
         open(resume_paths[1], "rb") as r2, \
         open(resume_paths[2], "rb") as r3:
        
        files = [
            ("jd_file", (os.path.basename(jd_path), jd_f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("resumes", (os.path.basename(resume_paths[0]), r1, "text/plain")),
            ("resumes", (os.path.basename(resume_paths[1]), r2, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("resumes", (os.path.basename(resume_paths[2]), r3, "text/plain")),
        ]
        data = {
            "title": "Senior Backend Engineer (Python / Flask)"
        }
        resp = requests.post(f"{base_url}/api/upload", data=data, files=files)
        print(f"    Upload Status: {resp.status_code}, Response: {resp.json()}")
        assert resp.status_code == 200
        upload_data = resp.json()
        job_id = upload_data["job_id"]

    print(f"[4] Testing Multi-Agent Pipeline (/api/analyze/{job_id})...")
    analyze_resp = requests.post(f"{base_url}/api/analyze/{job_id}")
    print(f"    Analyze Status: {analyze_resp.status_code}, Response: {analyze_resp.json()}")
    assert analyze_resp.status_code == 200

    print(f"[5] Testing Dashboard Route (/dashboard/{job_id})...")
    dash_resp = requests.get(f"{base_url}/dashboard/{job_id}")
    print(f"    Dashboard Status: {dash_resp.status_code}")
    assert dash_resp.status_code == 200
    assert "Alex Morgan" in dash_resp.text

    print(f"[6] Testing Export Route (/api/job/{job_id}/export)...")
    export_resp = requests.get(f"{base_url}/api/job/{job_id}/export")
    print(f"    Export Status: {export_resp.status_code}")
    export_json = export_resp.json()
    print(f"    Total Exported Candidates: {len(export_json['candidates'])}")
    for c in export_json['candidates']:
        print(f"      - {c['name']} (Score: {c['match_score']}, Tier: {c['match_tier']})")

    first_cand_id = 1
    print(f"[7] Testing Candidate Dossier Route (/candidate/{first_cand_id})...")
    cand_resp = requests.get(f"{base_url}/candidate/{first_cand_id}")
    print(f"    Candidate Dossier Status: {cand_resp.status_code}")
    assert cand_resp.status_code == 200

    print("\n[SUCCESS] ALL TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_test()
